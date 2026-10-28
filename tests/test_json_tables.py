"""Admin JSON data shown as a spreadsheet: shape detection and the Excel workbook (pure logic)."""
import io

import openpyxl

from app.services import jsontable
from app.services.export import build_workbook, sheet_titles


def only(data):
    sheets = jsontable.to_sheets(data)
    assert len(sheets) == 1
    return sheets[0]


# ---- shapes --------------------------------------------------------------------------------
def test_list_of_objects_is_one_row_per_object_with_the_union_of_keys():
    sheet = only([{"Д/Д": 1, "Овог": "Бат"}, {"Д/Д": 2, "Нас": 19}])
    assert sheet.columns == ["Д/Д", "Овог", "Нас"]                 # first-seen order
    assert sheet.rows == [[1, "Бат", None], [2, None, 19]]


def test_list_of_lists_is_a_plain_grid_with_spreadsheet_column_names():
    sheet = only([["Д/Д", "Овог"], [1, "Бат", "илүү"], [2]])
    assert sheet.columns == ["A", "B", "C"]
    assert sheet.rows == [["Д/Д", "Овог", None], [1, "Бат", "илүү"], [2, None, None]]


def test_list_of_plain_values_is_one_column():
    sheet = only([1, "a", None])
    assert sheet.columns == [jsontable.VALUE_COLUMN] and sheet.rows == [[1], ["a"], [None]]
    assert only([]).rows == []                                     # an empty list is an empty sheet, not an error


def test_mixed_list_does_not_crash_and_shows_nested_values_as_json_text():
    sheet = only([{"a": 1}, 5, [1, 2]])
    assert sheet.rows == [['{"a": 1}'], [5], ["[1, 2]"]]


def test_object_of_tables_gives_one_sheet_per_key_and_leftovers_go_to_an_extra_sheet():
    sheets = jsontable.to_sheets({
        "Сурагчид": [{"нэр": "Бат"}],
        "Багш": [["a", "b"]],
        "version": 3,
        "meta": {"by": "admin"},
    })
    assert [s.name for s in sheets] == ["Сурагчид", "Багш", jsontable.OTHER_SHEET]
    assert sheets[0].rows == [["Бат"]] and sheets[1].columns == ["A", "B"]
    assert sheets[2].rows == [["version", 3], ["meta", '{"by": "admin"}']]


def test_wrapper_object_with_a_single_list_is_that_table():
    sheets = jsontable.to_sheets({"data": [{"x": 1}, {"x": 2}]})
    assert len(sheets) == 1 and sheets[0].name == "data" and sheets[0].rows == [[1], [2]]


def test_object_of_scalar_lists_is_column_oriented_and_ragged_columns_are_padded():
    sheet = only({"Нэр": ["Бат", "Дорж"], "Нас": [19]})
    assert sheet.columns == ["Нэр", "Нас"] and sheet.rows == [["Бат", 19], ["Дорж", None]]


def test_plain_object_is_a_key_value_table():
    sheet = only({"a": 1, "b": {"c": 2}})
    assert sheet.columns == [jsontable.KEY_COLUMN, jsontable.VALUE_COLUMN]
    assert sheet.rows == [["a", 1], ["b", '{"c": 2}']]
    assert only({}).columns == [jsontable.KEY_COLUMN, jsontable.VALUE_COLUMN]


def test_a_bare_value_is_a_single_cell():
    assert only("hello").rows == [["hello"]] and only(42).rows == [[42]] and only(None).rows == [[None]]


# ---- cells ---------------------------------------------------------------------------------
def test_cell_values_keep_numbers_and_booleans_and_stringify_the_rest():
    assert jsontable.cell(1.5) == 1.5 and jsontable.cell(True) is True and jsontable.cell(None) is None
    assert jsontable.cell(float("nan")) == "nan"                   # NaN would corrupt a workbook
    assert jsontable.cell({"м": [1, 2]}) == '{"м": [1, 2]}'        # Cyrillic stays readable, not \u escapes


def test_column_letters():
    assert [jsontable.column_letters(i) for i in (0, 1, 25, 26, 27, 51, 52, 701, 702)] == \
        ["A", "B", "Z", "AA", "AB", "AZ", "BA", "ZZ", "AAA"]


def test_absurdly_wide_data_is_cut_not_exploded():
    sheet = only([{f"k{i}": i for i in range(jsontable.MAX_COLUMNS + 50)}])
    assert len(sheet.columns) == jsontable.MAX_COLUMNS and len(sheet.rows[0]) == jsontable.MAX_COLUMNS
    assert sheet.columns_cut is True


def test_parse_reads_json_text_and_reports_bad_input():
    assert jsontable.parse('[{"a": 1}]')[0].rows == [[1]]
    for bad in ("{not json", "[" * 100000):
        try:
            jsontable.parse(bad)
        except ValueError:
            pass
        else:
            raise AssertionError(f"accepted {bad[:10]!r}")


# ---- Excel workbook ------------------------------------------------------------------------
def test_sheet_titles_are_legal_unique_and_never_empty():
    titles = sheet_titles(["a/b:c", "A/B:C", "x" * 40, "x" * 40, "", "[ ]*?"])
    assert titles[0] == "a_b_c" and titles[1] == "A_B_C (2)"
    assert all(1 <= len(t) <= 31 and not set(t) & set("[]:*?/\\") for t in titles)
    assert len({t.casefold() for t in titles}) == len(titles)


def test_workbook_has_one_sheet_per_table_with_numbers_kept_as_numbers():
    sheets = jsontable.to_sheets({"Сурагчид": [{"нэр": "Бат", "нас": 19, "ok": True, "x": None}], "version": 3})
    book = build_workbook([(s.name, s.columns, s.rows) for s in sheets])
    wb = openpyxl.load_workbook(io.BytesIO(book.read()))
    assert wb.sheetnames == ["Сурагчид", "Бусад"]
    rows = list(wb["Сурагчид"].iter_rows(values_only=True))
    assert rows[0] == ("нэр", "нас", "ok", "x") and rows[1] == ("Бат", 19, True, None)


def test_workbook_survives_control_characters_and_blocks_formulas():
    data = [{"=HYPERLINK(\"http://evil\")": "@SUM(A1)", "t": "a\x00b\x0bc", "long": "я" * 40000}]
    sheet = only(data)
    book = build_workbook([(sheet.name, sheet.columns, sheet.rows)])         # must not raise
    rows = list(openpyxl.load_workbook(io.BytesIO(book.read())).active.iter_rows(values_only=True))
    assert rows[0][0].startswith("'=") and rows[1][0] == "'@SUM(A1)"          # formulas defused, headers too
    assert rows[1][1] == "abc"                                                 # control characters dropped
    assert len(rows[1][2]) == 32767                                            # Excel's cell limit
