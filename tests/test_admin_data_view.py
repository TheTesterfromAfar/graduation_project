"""Admin JSON data shown as a spreadsheet (and downloadable as Excel / CSV)."""
import csv
import io
import json

import openpyxl

from app.models import AdminDataEntry

from .conftest import add, first_id

ROSTER = [{"Д/Д": 1, "Овог": "Бат", "Нас": 19}, {"Д/Д": 2, "Овог": "Дорж", "Нас": 20}]
TWO_TABLES = {"Сурагчид": [{"нэр": "Бат"}], "Багш": [{"нэр": "Дорж"}], "version": 3}


def seed(app, data, filename="roster.json"):
    text = data if isinstance(data, str) else json.dumps(data, ensure_ascii=False)
    add(app, AdminDataEntry(filename=filename, content=text))
    return first_id(app, AdminDataEntry)


def page(client, url):
    response = client.get(url)
    assert response.status_code == 200, url
    return response.get_data(as_text=True)


def csv_rows(response):
    return list(csv.reader(io.StringIO(response.data.decode("utf-8-sig"))))


# ---- showing it ----------------------------------------------------------------------------
def test_pasted_json_is_shown_as_a_spreadsheet_right_after_adding(app, admin_client):
    response = admin_client.post("/admin/data/add", data={"text_content": json.dumps(ROSTER, ensure_ascii=False)})
    entry_id = first_id(app, AdminDataEntry)
    assert response.status_code == 302 and response.headers["Location"].endswith(f"/admin/data/{entry_id}")

    body = page(admin_client, f"/admin/data/{entry_id}")
    assert 'class="sheet"' in body                                       # a table, not a block of JSON
    assert "Овог" in body and "Бат" in body and "Дорж" in body
    assert "2 мөр · 3 багана" in body


def test_uploaded_json_file_is_shown_as_a_spreadsheet_too(app, admin_client):
    upload = {"json_file": (io.BytesIO(("\ufeff" + json.dumps(ROSTER, ensure_ascii=False)).encode("utf-8")), "roster.json")}
    response = admin_client.post("/admin/data/add", data=upload, content_type="multipart/form-data")
    entry_id = first_id(app, AdminDataEntry)
    assert response.status_code == 302 and response.headers["Location"].endswith(f"/admin/data/{entry_id}")
    assert "Дорж" in page(admin_client, f"/admin/data/{entry_id}")


def test_existing_entries_open_as_spreadsheets_and_the_list_links_to_them(app, admin_client, analyst_client):
    entry_id = seed(app, ROSTER)
    assert f"/admin/data/{entry_id}" in page(admin_client, "/admin/data")

    admin_page = page(admin_client, f"/admin/data/{entry_id}")
    analyst_page = page(analyst_client, f"/admin/data/{entry_id}")          # read-only staff can look
    assert "Бат" in analyst_page
    assert "JSON засах" in admin_page and "JSON засах" not in analyst_page


def test_each_table_in_the_json_becomes_its_own_sheet(app, admin_client):
    entry_id = seed(app, TWO_TABLES)
    body = page(admin_client, f"/admin/data/{entry_id}")
    assert 'id="sheet-0"' in body and 'id="sheet-1"' in body and 'id="sheet-2"' in body
    assert "Сурагчид" in body and "Багш" in body and "Бусад" in body        # leftovers: key/value sheet


def test_cell_text_cannot_inject_html(app, admin_client):
    entry_id = seed(app, [{"a": "<script>alert(1)</script>"}])
    body = page(admin_client, f"/admin/data/{entry_id}")
    assert "<script>alert(1)</script>" not in body and "&lt;script&gt;alert(1)&lt;/script&gt;" in body


def test_long_sheets_are_capped_on_screen_but_complete_in_the_download(app, admin_client):
    entry_id = seed(app, [{"n": i} for i in range(1500)])
    body = page(admin_client, f"/admin/data/{entry_id}")
    assert "Эхний 1000 мөрийг харуулж байна (нийт 1500)" in body
    sheet = openpyxl.load_workbook(io.BytesIO(admin_client.get(f"/admin/data/{entry_id}/export/xlsx").data)).active
    assert len(list(sheet.iter_rows())) == 1501                              # header + every row


def test_unreadable_stored_content_shows_a_message_instead_of_failing(app, admin_client):
    entry_id = seed(app, "this is not json")
    assert "боломжгүй" in page(admin_client, f"/admin/data/{entry_id}")
    assert admin_client.get(f"/admin/data/{entry_id}/export/xlsx").status_code == 302


# ---- downloading it ------------------------------------------------------------------------
def test_excel_download_has_one_worksheet_per_sheet(app, admin_client):
    entry_id = seed(app, TWO_TABLES)
    response = admin_client.get(f"/admin/data/{entry_id}/export/xlsx")
    assert response.status_code == 200 and "spreadsheetml" in response.mimetype
    workbook = openpyxl.load_workbook(io.BytesIO(response.data))
    assert workbook.sheetnames == ["Сурагчид", "Багш", "Бусад"]
    assert list(workbook["Сурагчид"].iter_rows(values_only=True)) == [("нэр",), ("Бат",)]


def test_csv_download_is_excel_friendly_formula_safe_and_picks_a_sheet(app, admin_client):
    entry_id = seed(app, {"Эхний": [{"a": "=1+1"}], "Хоёр": [{"b": 2}]})
    first = admin_client.get(f"/admin/data/{entry_id}/export/csv")
    assert first.mimetype == "text/csv" and first.data.startswith(b"\xef\xbb\xbf")   # UTF-8 BOM for Excel
    assert csv_rows(first) == [["a"], ["'=1+1"]]                                    # formula defused
    assert csv_rows(admin_client.get(f"/admin/data/{entry_id}/export/csv?sheet=1")) == [["b"], ["2"]]
    assert admin_client.get(f"/admin/data/{entry_id}/export/csv?sheet=2").status_code == 404


def test_unknown_entry_or_format_is_404(app, admin_client):
    entry_id = seed(app, ROSTER)
    assert admin_client.get("/admin/data/999").status_code == 404
    assert admin_client.get("/admin/data/999/export/csv").status_code == 404
    assert admin_client.get(f"/admin/data/{entry_id}/export/exe").status_code == 404


def test_login_is_required(app, client):
    entry_id = seed(app, ROSTER)
    assert client.get(f"/admin/data/{entry_id}").status_code == 302
    assert client.get(f"/admin/data/{entry_id}/export/xlsx").status_code == 302


# ---- editing -------------------------------------------------------------------------------
def test_saving_an_edit_returns_to_the_spreadsheet(app, admin_client):
    entry_id = seed(app, ROSTER)
    response = admin_client.post(f"/admin/data/{entry_id}/edit", data={"content": '[{"Нэр": "Шинэ нэр"}]'})
    assert response.status_code == 302 and response.headers["Location"].endswith(f"/admin/data/{entry_id}")
    body = page(admin_client, f"/admin/data/{entry_id}")
    assert "Шинэ нэр" in body and "Дорж" not in body
